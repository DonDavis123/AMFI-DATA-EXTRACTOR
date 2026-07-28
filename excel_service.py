import os

from openpyxl import Workbook, load_workbook


class ExcelService:

    def __init__(self):

        self.output_folder = "output"
        self.file_path = os.path.join(self.output_folder, "funds.xlsx")

        os.makedirs(self.output_folder, exist_ok=True)

        # Create workbook if it doesn't exist
        if not os.path.exists(self.file_path):

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Mutual Funds"

            sheet.append([
                "Fund Name",
                "ISIN",
                "Fund Type",
                "Riskometer At Launch",
                "Riskometer As On Date",
                "Category",
                "Description",
                "Fund Manager Name"
            ])

            workbook.save(self.file_path)
            workbook.close()

    def append(self, extraction):
        """
        extraction -> FundExtraction
        """

        try:

            # Always load the latest workbook from disk
            workbook = load_workbook(self.file_path)
            sheet = workbook.active

            for fund in extraction.funds:

                sheet.append([
                    fund.fund_name,
                    fund.isin,
                    fund.fund_type,
                    fund.riskometer_at_launch,
                    fund.riskometer_as_on_date,
                    fund.category,
                    fund.description,
                    fund.fund_manager_name
                ])

            workbook.save(self.file_path)
            workbook.close()

            print(f"✓ Saved {len(extraction.funds)} row(s) to Excel. - excel_service.py:63")

        except PermissionError:

            print("\n - excel_service.py:67" + "=" * 80)
            print("ERROR - excel_service.py:68")
            print("The Excel file 'funds.xlsx' is currently open. - excel_service.py:69")
            print("Please close the file and run the program again. - excel_service.py:70")
            print("= - excel_service.py:71" * 80)

            raise

        except Exception as e:

            print("\n - excel_service.py:77" + "=" * 80)
            print(f"Failed to save Excel file: {e} - excel_service.py:78")
            print("= - excel_service.py:79" * 80)

            raise